"""ML2-6: регрессия совместимости выходного PPTX с Р7 (см. docs/R7_COMPAT_CHECKLIST.md)."""
import shutil
import subprocess
from pathlib import Path

import pytest
from pptx import Presentation

from deckgen.contracts import DeckPlan, SlideSpec, SlideType, TableSpec
from deckgen.fitter import DeterministicFitter
from deckgen.renderer import PptxRenderer
from deckgen.renderer.sanity import check_pptx
from deckgen.template import profile

import ml2_fixtures as fx

ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = {"example": ROOT / "templates" / "example_fields.pptx",
             "interrao": ROOT / "templates" / "interrao_corporate.pptx"}


def _render(tpl: Path, out: Path, plan=None) -> Path:
    cat = profile(str(tpl))
    deck = DeterministicFitter().fit(plan or fx.plan(), fx.ingest(), cat)
    PptxRenderer().render(deck, str(tpl), cat, str(out))
    return out


@pytest.mark.parametrize("key", list(TEMPLATES))
def test_output_is_clean(key, tmp_path):
    out = _render(TEMPLATES[key], tmp_path / "o.pptx")
    problems = check_pptx(str(out))
    assert not problems, "\n".join(map(str, problems))
    prs = Presentation(str(out))  # открывается заново
    assert prs.core_properties.title == fx.plan().slides[0].title


def test_checker_catches_template_garbage():
    # исходный шаблон: autofit в макетах, Calibri в теме — чекер это видит
    codes = {p.code for p in check_pptx(str(TEMPLATES["example"]))}
    assert {"AUTOFIT", "FONT"} <= codes


def test_checker_invalid_file(tmp_path):
    bad = tmp_path / "x.pptx"
    bad.write_bytes(b"nope")
    assert check_pptx(str(bad))[0].code == "INVALID"


def _foreign_template(tmp_path) -> Path:
    """Шаблон «как у жюри»: те же поля, но другие имена макетов, нет table_only* и last*."""
    prs = Presentation(str(TEMPLATES["example"]))
    rename = {"title1": "Title_main", "section1": "section_blue", "slide1": "Slide_text",
              "slide2": "slide_text_image", "table_text1": "table_with_text"}
    layouts = prs.slide_masters[0].slide_layouts
    for l in list(layouts):
        if l.name in ("table_only1", "last1"):
            layouts.remove(l)
        elif l.name in rename:
            l.name = rename[l.name]
    out = tmp_path / "jury.pptx"
    prs.save(str(out))
    return out


def test_foreign_convention_template(tmp_path):
    tpl = _foreign_template(tmp_path)
    cat = profile(str(tpl))
    assert cat.adapter == "convention"
    assert {l.layout_id for l in cat.layouts} == {"Title_main", "section_blue", "Slide_text",
                                                  "slide_text_image", "table_with_text"}
    out = _render(tpl, tmp_path / "o.pptx")
    assert not check_pptx(str(out))
    deck = DeterministicFitter().fit(fx.plan(), fx.ingest(), cat)
    by_type = {s.spec.type: s.layout_id for s in deck.slides}
    assert by_type[SlideType.last] == "section_blue"          # нет last* -> фолбэк на раздел
    assert all(s.layout_id == "table_with_text" for s in deck.slides if s.spec.table)


def test_edge_plans_render_clean(tmp_path):
    plan = DeckPlan(deck_title="", slides=[
        SlideSpec(type=SlideType.title, title="Очень длинное название презентации про развитие "
                  "распределительного сетевого комплекса региона на период до 2030 года"),
        SlideSpec(type=SlideType.slide, title="Спецсимволы <>&\"'", bullets=["A & B < C > D", "• уже с маркером"]),
        SlideSpec(type=SlideType.table, title="Таблица без источника",
                  table=TableSpec(columns=["a", "b"], rows=[["1", "2"], ["3"]])),
        SlideSpec(type=SlideType.slide, title="Только подзаголовок", subtitle="Пояснение без буллетов"),
    ])
    for key, tpl in TEMPLATES.items():
        out = _render(tpl, tmp_path / f"{key}.pptx", plan)
        assert not check_pptx(str(out)), key


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice не установлен")
@pytest.mark.parametrize("key", list(TEMPLATES))
def test_libreoffice_converts(key, tmp_path):
    out = _render(TEMPLATES[key], tmp_path / "o.pptx")
    r = subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp_path), str(out)],
                       capture_output=True, text=True, timeout=180)
    assert r.returncode == 0 and (tmp_path / "o.pdf").exists(), r.stderr


def test_unfilled_sample_text_removed(tmp_path):
    """ИнтерРАО: титул без подзаголовка — текст образца «Докладчик: ФИО» не должен остаться."""
    plan = DeckPlan(deck_title="x", slides=[SlideSpec(type=SlideType.title, title="Отчёт")])
    out = _render(TEMPLATES["interrao"], tmp_path / "o.pptx", plan)
    text = " ".join(sh.text_frame.text for sh in Presentation(str(out)).slides[0].shapes if sh.has_text_frame)
    assert "ФИО" not in text and "Отчёт" in text
