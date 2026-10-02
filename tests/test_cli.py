"""ML1-4: CLI сквозняком на заглушках (или реальных модулях, если они уже есть). Сеть не нужна."""
from __future__ import annotations
import json

from pptx import Presentation
from typer.testing import CliRunner

from deckgen.cli import app
from deckgen.contracts import DeckPlan, ValidationReport

runner = CliRunner()


def _check_outputs(out):
    assert out.exists()
    plan = DeckPlan.model_validate_json(out.with_suffix(".plan.json").read_text(encoding="utf-8"))
    report = ValidationReport.model_validate_json(out.with_suffix(".report.json").read_text(encoding="utf-8"))
    prs = Presentation(str(out))
    assert len(prs.slides) >= len(plan.slides) > 2
    return plan, report, prs


def test_demo_no_llm(tmp_path):
    out = tmp_path / "demo.pptx"
    r = runner.invoke(app, ["--demo", "--no-llm", "-o", str(out), "--env-file", str(tmp_path / "none.env")])
    assert r.exit_code == 0, r.output
    plan, report, prs = _check_outputs(out)
    assert "fit" in r.output and "render" in r.output and "Отчёт" in r.output
    assert report.metrics["fallback_used"] == 0  # --no-llm — это не авария, а явный режим
    # таблица в PPTX — нативная, с ячейками ровно как в источнике
    tables = [sh.table for s in prs.slides for sh in s.shapes if sh.has_table]
    assert tables and tables[0].cell(1, 2).text == "18,7" and tables[0].cell(5, 2).text == "≤ 5"


def test_demo_without_network_falls_back(tmp_path, monkeypatch):
    monkeypatch.setenv("LLM_BASE_URL", "http://127.0.0.1:9/v1")   # закрытый порт = «сеть выключена»
    monkeypatch.setenv("LLM_MAX_RETRIES", "0")
    monkeypatch.setenv("LLM_LOG_PATH", str(tmp_path / "llm.jsonl"))
    out = tmp_path / "o.pptx"
    r = runner.invoke(app, ["--demo", "-o", str(out), "--env-file", str(tmp_path / "none.env")])
    assert r.exit_code == 0, r.output
    _, report, _ = _check_outputs(out)
    assert "fallback" in r.output and report.metrics["fallback_used"] == 1


def test_input_file(tmp_path):
    src = tmp_path / "in.md"
    src.write_text("# Совещание\n\n## Поручения\n- Подготовить отчёт до 15 марта 2025 года.\n- Провести 2 проверки.\n",
                   encoding="utf-8")
    out = tmp_path / "r.pptx"
    r = runner.invoke(app, ["--input", str(src), "--no-llm", "-o", str(out), "--env-file", str(tmp_path / "x")])
    assert r.exit_code == 0, r.output
    plan, report, _ = _check_outputs(out)
    assert report.ok, report.issues
    assert any("15 марта 2025" in e.quote for s in plan.slides for e in s.evidence)


def test_requires_input():
    r = runner.invoke(app, [])
    assert r.exit_code != 0
