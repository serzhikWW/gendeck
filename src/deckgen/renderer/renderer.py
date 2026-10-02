"""Renderer: FittedDeck + шаблон -> редактируемый PPTX.

convention: слайд = add_slide(макет), поля = плейсхолдеры по idx, таблица = insert_table в TABLE_FIELD.
profile:    слайд = клон слайда-образца, поля = фигуры по имени из YAML или новые TextBox в зонах,
            таблица = свой graphicFrame в зоне TABLE.
Потом: удалить слайды-образцы шаблона, пустые плейсхолдеры, autofit; перенумеровать части.
"""
from __future__ import annotations

import copy
from pathlib import Path
from typing import Optional

from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Emu

from deckgen.contracts import BBox, FieldKind, FieldSpec, FittedDeck, FittedSlide, LayoutSpec, SlideType, TemplateCatalog
from deckgen.template.convention import find_layout, open_presentation
from deckgen.template.geometry import content_zones
from deckgen.template.profile import Profile, find_profile
from deckgen.template.textstyle import (BULLET_CHAR, BULLET_MARL_EMU, LN_SPC, PARA_SPACE_RATIO,
                                        TEXT_INSETS_EMU)

from . import ops
from .tables import TableStyle, fill_table

_CONTENT = (SlideType.slide, SlideType.table)
_DEFAULT_BULLET = "F08800"


class PptxRenderer:
    def render(self, deck: FittedDeck, template_path: str, catalog: TemplateCatalog, out_path: str) -> None:
        prs = open_presentation(template_path)
        originals = list(prs.slides)
        layouts = {l.layout_id: l for l in catalog.layouts}
        profile = find_profile(template_path, prs) if catalog.adapter == "profile" else None
        ctx = _Ctx(prs, catalog, profile)
        for fs in deck.slides:
            lspec = layouts.get(fs.layout_id)
            if lspec is None:
                raise ValueError(f"layout_id '{fs.layout_id}' нет в каталоге {catalog.template_id}")
            if lspec.source == "sample_slide":
                slide = ops.clone_slide(prs, originals[lspec.sample_slide_index])
            else:
                slide = prs.slides.add_slide(find_layout(prs, lspec.layout_id))
            _SlideWriter(ctx, slide, fs, lspec).write()
        ops.delete_slides(prs, originals)
        for s in prs.slides:
            ops.remove_empty_placeholders(s)
        ops.strip_autofit_everywhere(prs)
        ops.renumber(prs)
        first_title = next((s.spec.title for s in deck.slides if s.spec.type == SlideType.title), None)
        if first_title:
            prs.core_properties.title = first_title[:255]
        Path(out_path).parent.mkdir(parents=True, exist_ok=True)
        prs.save(out_path)


class _Ctx:
    def __init__(self, prs, catalog: TemplateCatalog, profile: Optional[Profile]):
        self.prs, self.catalog, self.profile = prs, catalog, profile
        self._style_cache: dict[str, TableStyle] = {}

    def color(self, key: str, default: Optional[str] = None) -> Optional[str]:
        if self.profile:
            v = self.profile.colors.get(key)
            return v or default
        return default

    def table_style(self, lspec: LayoutSpec) -> TableStyle:
        if lspec.layout_id in self._style_cache:
            return self._style_cache[lspec.layout_id]
        if self.profile:
            st = TableStyle(style_id="{2D5ABB26-0587-4C30-8999-92F81FD0307C}",  # No Style, No Grid
                            header_text=self.color("table_header_text"), header_fill=self.color("table_header_fill"),
                            band_fill=self.color("table_band_fill"), border=self.color("table_border", "D9D9D9"))
        else:
            st = TableStyle(style_id=_template_table_style(self.prs, lspec.layout_id))
        self._style_cache[lspec.layout_id] = st
        return st


def _template_table_style(prs, layout_id: str) -> Optional[str]:
    """tableStyleId: таблица STYLE_TABLE (или любая) в макете -> в мастере -> def из tableStyles.xml."""
    layout = find_layout(prs, layout_id)
    for root in (layout._element, layout.slide_master._element):
        for gf in root.iter(qn("p:graphicFrame")):
            sid = gf.find(f".//{qn('a:tableStyleId')}")
            if sid is not None and sid.text:
                return sid.text.strip()
    try:
        ts = prs.part.part_related_by("http://schemas.openxmlformats.org/officeDocument/2006/relationships/tableStyles")
        root = etree.fromstring(ts.blob)
        return root.get("def")
    except Exception:
        return None


def _layout_bullet_color(slide) -> Optional[str]:
    """Цвет буллета из TEXT-плейсхолдера макета (buClr), чтобы буллеты были «как в шаблоне»."""
    for sp in slide.slide_layout.placeholders:
        if sp.placeholder_format.type is not None and "BODY" in str(sp.placeholder_format.type):
            for clr in sp._element.iter(qn("a:buClr")):
                s = clr.find(qn("a:srgbClr"))
                if s is not None:
                    return s.get("val")
    return None


class _SlideWriter:
    def __init__(self, ctx: _Ctx, slide, fs: FittedSlide, lspec: LayoutSpec):
        self.ctx, self.slide, self.fs, self.lspec = ctx, slide, fs, lspec
        self.spec = fs.spec

    # ---- поиск фигур ----
    def _placeholder(self, idx: Optional[int]):
        if idx is None:
            return None
        for ph in self.slide.placeholders:
            if ph.placeholder_format.idx == idx:
                return ph
        return None

    def _shape_by_name(self, name: str):
        for sh in self.slide.shapes:
            if sh.name == name:
                return sh
        return None

    def _profile_field(self, f: FieldSpec) -> dict:
        if not self.ctx.profile:
            return {}
        try:
            return self.ctx.profile.field(self.lspec.layout_id, f.name)
        except Exception:
            return {}

    def _target(self, f: FieldSpec, bbox: BBox):
        """Фигура для поля: плейсхолдер по idx -> фигура образца по имени -> новый TextBox."""
        sh = self._placeholder(f.placeholder_idx)
        if sh is None:
            pf = self._profile_field(f)
            if pf.get("shape"):
                sh = self._shape_by_name(pf["shape"])
        created = False
        if sh is None:
            sh = self.slide.shapes.add_textbox(Emu(bbox.x), Emu(bbox.y), Emu(bbox.w), Emu(bbox.h))
            sh.name = f.name
            created = True
        if created or (sh.left, sh.top, sh.width, sh.height) != (bbox.x, bbox.y, bbox.w, bbox.h):
            sh.left, sh.top, sh.width, sh.height = Emu(bbox.x), Emu(bbox.y), Emu(bbox.w), Emu(bbox.h)
        return sh, created

    def font(self, f: FieldSpec, default: float) -> float:
        return float(self.fs.font_pt.get(f.name) or f.font_pt or default)

    # ---- запись ----
    def write(self) -> None:
        st = self.ctx.catalog.style
        subtitle_written = False
        for f in self.lspec.fields:
            if f.kind == FieldKind.title and self.spec.title:
                self._write_plain(f, self.spec.title, self.font(f, st.title_pt), title=True)
            elif f.kind == FieldKind.subtitle and self.spec.subtitle:
                self._write_plain(f, self.spec.subtitle, self.font(f, st.body_pt_max), title=False)
                subtitle_written = True
        has_table = self.spec.table is not None
        lead = self.spec.subtitle if not subtitle_written else ""
        paras_exist = bool(self.spec.bullets or lead)
        zones = content_zones(self.lspec, paras_exist, has_table)
        if paras_exist and "text" in zones:
            z = zones["text"]
            self._write_bullets(z, lead, self.spec.bullets, self.font(z, st.body_pt_max))
        if has_table and "table" in zones:
            z = zones["table"]
            self._write_table(z, self.font(z, st.table_pt_max))
        if self.spec.notes:
            self.slide.notes_slide.notes_text_frame.text = self.spec.notes

    def _write_plain(self, f: FieldSpec, text: str, size: float, title: bool) -> None:
        sh, created = self._target(f, f.bbox)
        tf = sh.text_frame
        tmpl_rpr, tmpl_ppr = _first_rpr_ppr(sh._element)
        tf.text = ""
        tf.word_wrap = True
        p = tf.paragraphs[0]
        if tmpl_ppr is not None:
            for old in p._p.findall(qn("a:pPr")):
                p._p.remove(old)
            p._p.insert(0, copy.deepcopy(tmpl_ppr))
        lines = text.split("\n")
        for i, line in enumerate(lines):
            if i:
                p.add_line_break()
            r = p.add_run()
            r.text = line
            if tmpl_rpr is not None:
                _apply_rpr(r, tmpl_rpr)
            color = None
            if created:
                color = self.ctx.color("title" if title else "text")
            ops.style_run(r, size, color=color)
        if created or (title and tmpl_ppr is None):
            # выравнивание как у подсказки заголовка в макете (её pPr не наследуется, но это замысел автора)
            p.alignment = _title_alignment(self.slide, self.lspec)
        ops.end_para_size(p, size)
        ops.no_autofit(sh._element.find(qn("p:txBody")))

    def _write_bullets(self, z: FieldSpec, lead: str, bullets: list[str], size: float) -> None:
        sh, created = self._target(z, z.bbox)
        if created:
            l, t, r, b = TEXT_INSETS_EMU
            tf = sh.text_frame
            tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = Emu(l), Emu(t), Emu(r), Emu(b)
        tf = sh.text_frame
        tf.text = ""
        tf.word_wrap = True
        bullet_color = self.ctx.color("accent") if self.ctx.profile else (_layout_bullet_color(self.slide) or _DEFAULT_BULLET)
        text_color = self.ctx.color("text") if created else None
        items = ([(lead, False)] if lead else []) + [(b, True) for b in bullets]
        for i, (txt, is_bullet) in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            _bullet_ppr(p, size, is_bullet, bullet_color, first=(i == 0))
            for j, line in enumerate(txt.split("\n")):
                if j:
                    p.add_line_break()
                r = p.add_run()
                r.text = line
                ops.style_run(r, size, bold=None if is_bullet else True, color=text_color)
            ops.end_para_size(p, size)
        ops.no_autofit(sh._element.find(qn("p:txBody")))

    def _write_table(self, z: FieldSpec, size: float) -> None:
        from deckgen.fitter.measure import table_grid
        grid, _ = table_grid(self.spec.table)
        rows, cols = len(grid), len(grid[0]) if grid else 1
        if rows == 0:
            return
        ph = self._placeholder(z.placeholder_idx)
        if ph is not None and hasattr(ph, "insert_table") and z.kind == FieldKind.table:
            gf = ph.insert_table(rows, cols)
        else:
            gf = self.slide.shapes.add_table(rows, cols, Emu(z.bbox.x), Emu(z.bbox.y), Emu(z.bbox.w), Emu(z.bbox.h))
        gf.name = "TABLE"
        fill_table(gf, self.spec.table, z.bbox, size, self.ctx.table_style(self.lspec))


# ---------- XML-помощники ----------
def _first_rpr_ppr(sp):
    """rPr/pPr первого рана фигуры-образца (чтобы сохранить её цвет/жирность/выравнивание)."""
    txb = sp.find(qn("p:txBody"))
    if txb is None:
        return None, None
    rpr = None
    for r in txb.iter(qn("a:r")):
        rpr = r.find(qn("a:rPr"))
        break
    p = txb.find(qn("a:p"))
    ppr = p.find(qn("a:pPr")) if p is not None else None
    return rpr, ppr


def _apply_rpr(run, tmpl_rpr) -> None:
    new = copy.deepcopy(tmpl_rpr)
    for attr in ("sz", "dirty", "err", "smtClean"):
        new.attrib.pop(attr, None)
    old = run._r.find(qn("a:rPr"))
    if old is not None:
        run._r.remove(old)
    run._r.insert(0, new)


def _title_alignment(slide, lspec: LayoutSpec):
    """Выравнивание как у заголовка: lvl1pPr@algn плейсхолдера заголовка макета -> titleStyle мастера."""
    from pptx.enum.text import PP_ALIGN
    algn_map = {"l": PP_ALIGN.LEFT, "ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "just": PP_ALIGN.JUSTIFY}
    layout = slide.slide_layout
    roots = [sp._element for sp in layout.placeholders
             if sp.placeholder_format.idx == 0]
    ts = layout.slide_master._element.find(f"{qn('p:txStyles')}/{qn('p:titleStyle')}")
    if ts is not None:
        roots.append(ts)
    for root in roots:
        for lvl in root.iter(qn("a:lvl1pPr")):
            if lvl.get("algn") in algn_map:
                return algn_map[lvl.get("algn")]
        for ppr in root.iter(qn("a:pPr")):
            if ppr.get("algn") in algn_map:
                return algn_map[ppr.get("algn")]
    return PP_ALIGN.CENTER if lspec.slide_type in (SlideType.title, SlideType.section, SlideType.last) else PP_ALIGN.LEFT


def _bullet_ppr(p, size: float, is_bullet: bool, color: Optional[str], first: bool) -> None:
    ppr = p._p.get_or_add_pPr()
    for child in list(ppr):
        ppr.remove(child)
    if is_bullet:
        ppr.set("marL", str(BULLET_MARL_EMU))
        ppr.set("indent", str(-BULLET_MARL_EMU))
    else:
        ppr.set("marL", "0")
        ppr.set("indent", "0")
    ppr.set("lvl", "0")
    ppr.set("algn", "l")
    ln = etree.SubElement(ppr, qn("a:lnSpc"))
    etree.SubElement(ln, qn("a:spcPct"), val=str(int(LN_SPC * 100000)))
    bef = etree.SubElement(ppr, qn("a:spcBef"))
    etree.SubElement(bef, qn("a:spcPts"), val="0" if first else str(int(round(PARA_SPACE_RATIO * size * 100))))
    aft = etree.SubElement(ppr, qn("a:spcAft"))
    etree.SubElement(aft, qn("a:spcPts"), val="0")
    if is_bullet:
        if color:
            bc = etree.SubElement(ppr, qn("a:buClr"))
            etree.SubElement(bc, qn("a:srgbClr"), val=color)
        etree.SubElement(ppr, qn("a:buFont"), typeface="Arial", panose="020B0604020202020204",
                         pitchFamily="34", charset="0")
        etree.SubElement(ppr, qn("a:buChar"), char=BULLET_CHAR)
    else:
        etree.SubElement(ppr, qn("a:buNone"))
