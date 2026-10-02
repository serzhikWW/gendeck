"""ВРЕМЕННЫЕ заглушки модулей других ролей (владелец файла: ML1). Нужны только чтобы `cli --demo` и сквозной
путь работали, пока нет реальных реализаций. pipeline.default_modules() берёт реальный модуль, как только его
пакет экспортирует create() (deckgen.template/ingest/fitter/renderer/validator), — эти классы тогда не
используются. Минимальны сознательно: без разбиения слайдов, без оценки переполнения, только convention-шаблон.
"""
from __future__ import annotations
import hashlib
import re
from typing import Optional

from .contracts import *
from .planner.numbers import number_in_source, numbers_in, quote_in_source

FONT = "Liberation Sans"
_PREFIX = {"title": SlideType.title, "section": SlideType.section, "slide": SlideType.slide,
           "table": SlideType.table, "last": SlideType.last}
_KIND = {"TITLE_FIELD": FieldKind.title, "SUBTITLE_FIELD": FieldKind.subtitle, "TEXT_FIELD": FieldKind.text,
         "TABLE_FIELD": FieldKind.table, "IMAGE_FIELD": FieldKind.image}
_EMU_PT = 12700


def _slide_type(name: str) -> Optional[SlideType]:
    for p, t in _PREFIX.items():
        if name.lower().startswith(p):
            return t
    return None


class StubProfiler:
    """Convention-шаблон: макеты title*/section*/slide*/table*/last*, плейсхолдеры *_FIELD."""
    def profile(self, template_path: str) -> TemplateCatalog:
        from pptx import Presentation
        prs = Presentation(template_path)
        layouts = []
        for lay in prs.slide_layouts:
            st = _slide_type(lay.name)
            if st is None:
                continue
            fields = []
            for ph in lay.placeholders:
                kind = _KIND.get(ph.name)
                if kind is None:
                    continue
                font = 28.0 if kind == FieldKind.title and st in (SlideType.title, SlideType.section,
                                                                   SlideType.last) else \
                    24.0 if kind == FieldKind.title else 14.0
                w_pt, h_pt = ph.width / _EMU_PT, ph.height / _EMU_PT
                lines = max(1, int(h_pt / (font * 1.2)))
                per_line = max(1, int(w_pt / (font * 0.55)))
                fields.append(FieldSpec(name=ph.name, kind=kind, bbox=BBox(x=ph.left, y=ph.top, w=ph.width, h=ph.height),
                                        font_pt=font, max_chars=per_line * lines if kind != FieldKind.table else 0,
                                        max_lines=lines if kind != FieldKind.table else 0,
                                        placeholder_idx=ph.placeholder_format.idx))
            layouts.append(LayoutSpec(layout_id=lay.name, slide_type=st, fields=fields))
        if not layouts:
            raise NotImplementedError("Шаблон не в формате README (convention). Profile-режим — ML2-2.")
        return TemplateCatalog(template_id=template_path, adapter="convention", slide_w=prs.slide_width,
                               slide_h=prs.slide_height, layouts=layouts)


class StubIngestor:
    """Markdown: заголовки # -> секции, |-таблицы -> SourceTable (ячейки как в тексте)."""
    _SEP = re.compile(r"^\|?\s*:?-{2,}")

    def ingest(self, text: str) -> IngestResult:
        sections: list[SourceSection] = []
        tables: list[SourceTable] = []
        cur = SourceSection(section_id="s1")
        body: list[str] = []
        tbl: list[list[str]] = []
        prev_line = ""

        def flush_table():
            nonlocal tbl
            if len(tbl) >= 1:
                cap = prev_line if prev_line.lower().startswith("таблица") else cur.heading
                tid = f"t{len(tables) + 1}"
                tables.append(SourceTable(table_id=tid, columns=tbl[0], rows=tbl[1:], caption=cap))
                cur.table_ids.append(tid)
            tbl = []

        def flush_section():
            nonlocal cur, body
            cur.text = "\n".join(body).strip()
            if cur.heading or cur.text or cur.table_ids:
                sections.append(cur)
            cur = SourceSection(section_id=f"s{len(sections) + 1}")
            body = []

        for line in text.splitlines():
            s = line.strip()
            if s.startswith("|"):
                if not self._SEP.match(s):
                    tbl.append([c.strip() for c in s.strip("|").split("|")])
                continue
            if tbl:
                flush_table()
            m = re.match(r"^#{1,6}\s+(.*)", s)
            if m:
                flush_section()
                cur.heading = m.group(1).strip()
                continue
            body.append(line)
            if s:
                prev_line = s
        if tbl:
            flush_table()
        flush_section()
        return IngestResult(raw_text=text, source_hash=hashlib.sha256(text.encode("utf-8")).hexdigest(),
                            sections=sections, tables=tables)


class StubFitter:
    """Выбор макета по типу и наполнению; ячейки таблиц — строго из SourceTable. Без разбиения/сокращений."""
    def fit(self, plan: DeckPlan, src: IngestResult, catalog: TemplateCatalog) -> FittedDeck:
        tables = {t.table_id: t for t in src.tables}
        out = []
        for i, s in enumerate(plan.slides):
            s = s.model_copy(deep=True)
            if s.table and s.table.source_id in tables:
                t = tables[s.table.source_id]
                s.table.columns, s.table.rows = list(t.columns), [list(r) for r in t.rows]
            cands = catalog.by_type(s.type) or catalog.by_type(SlideType.slide)
            def has(l, k): return any(f.kind == k for f in l.fields)
            if s.type == SlideType.table:
                pref = [l for l in cands if has(l, FieldKind.text) == bool(s.bullets)] or cands
            else:
                pref = [l for l in cands if not has(l, FieldKind.image)] or cands
            out.append(FittedSlide(spec=s, layout_id=pref[0].layout_id, origin_slide_index=i))
        return FittedDeck(template_id=catalog.template_id, slides=out)


class StubRenderer:
    """python-pptx: слайды из макетов, явные кегли и Liberation Sans, нативные таблицы, без autofit."""
    def render(self, deck: FittedDeck, template_path: str, catalog: TemplateCatalog, out_path: str) -> None:
        from pathlib import Path
        from pptx import Presentation
        from pptx.enum.text import MSO_AUTO_SIZE
        from pptx.util import Pt

        prs = Presentation(template_path)
        sld_ids = prs.slides._sldIdLst
        for sld in list(sld_ids):                         # убрать слайды-образцы шаблона
            prs.part.drop_rel(sld.rId)
            sld_ids.remove(sld)
        layouts = {l.name: l for l in prs.slide_layouts}
        specs = {l.layout_id: l for l in catalog.layouts}

        def fill(tf, lines: list[str], size: float, bold=False):
            tf.clear()
            tf.word_wrap = True
            tf.auto_size = MSO_AUTO_SIZE.NONE
            for k, text in enumerate(lines):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                r = p.add_run()
                r.text = text
                r.font.size, r.font.name, r.font.bold = Pt(size), FONT, bold

        for fs in deck.slides:
            s = fs.spec
            lay = layouts[fs.layout_id]
            slide = prs.slides.add_slide(lay)
            by_idx = {f.placeholder_idx: f for f in specs[fs.layout_id].fields}
            for ph in list(slide.placeholders):
                f = by_idx.get(ph.placeholder_format.idx)
                kind = f.kind if f else None
                if kind == FieldKind.title and s.title:
                    fill(ph.text_frame, [s.title], f.font_pt, bold=True)
                elif kind == FieldKind.subtitle and s.subtitle:
                    fill(ph.text_frame, [s.subtitle], 18)
                elif kind == FieldKind.text and s.bullets:
                    fill(ph.text_frame, s.bullets, 14 if len(s.bullets) <= 5 else 12)
                elif kind == FieldKind.table and s.table and s.table.columns:
                    rows = [s.table.columns] + s.table.rows
                    gf = ph.insert_table(rows=len(rows), cols=len(s.table.columns))
                    for r, row in enumerate(rows):
                        for c, val in enumerate(row[:len(s.table.columns)]):
                            fill(gf.table.cell(r, c).text_frame, [val], 11 if len(rows) <= 12 else 9, bold=r == 0)
                else:
                    ph._element.getparent().remove(ph._element)   # пустые плейсхолдеры удаляются
            if s.subtitle and not any(f.kind == FieldKind.subtitle for f in specs[fs.layout_id].fields):
                from pptx.util import Emu
                tb = slide.shapes.add_textbox(Emu(int(prs.slide_width * 0.08)), Emu(int(prs.slide_height * 0.62)),
                                              Emu(int(prs.slide_width * 0.84)), Emu(int(prs.slide_height * 0.12)))
                fill(tb.text_frame, [s.subtitle], 18)
            if s.notes:
                slide.notes_slide.notes_text_frame.text = s.notes
        Path(out_path).parent.mkdir(parents=True, exist_ok=True)
        prs.save(out_path)


class StubValidator:
    """Мини-fidelity: числа слайдов есть в источнике, цитаты дословны, ячейки таблиц == SourceTable."""
    def validate(self, src: IngestResult, deck: FittedDeck, pptx_path: str, catalog: TemplateCatalog) -> ValidationReport:
        issues: list[Issue] = []
        n_total = n_ok = cells = cells_ok = 0
        tables = {t.table_id: t for t in src.tables}
        for i, fs in enumerate(deck.slides):
            s = fs.spec
            for text in [s.title, s.subtitle, *s.bullets]:
                for n in numbers_in(text):
                    n_total += 1
                    if number_in_source(n, src.raw_text):
                        n_ok += 1
                    else:
                        issues.append(Issue(severity="error", code=IssueCode.NUMBER_MISSING, slide_index=i,
                                            message=f"число {n} не найдено в источнике"))
            for e in s.evidence:
                if not quote_in_source(e.quote, src.raw_text):
                    issues.append(Issue(severity="error", code=IssueCode.EVIDENCE_NOT_FOUND, slide_index=i,
                                        message=f"цитата не найдена: «{e.quote[:60]}»"))
            if s.table and s.table.source_id in tables:
                t = tables[s.table.source_id]
                bad = 0
                for got, exp in zip([s.table.columns] + s.table.rows, [t.columns] + t.rows):
                    for g, x in zip(got, exp):
                        cells += 1
                        cells_ok += g == x
                        bad += g != x
                if bad:
                    issues.append(Issue(severity="error", code=IssueCode.TABLE_MISMATCH, slide_index=i,
                                        message=f"таблица {t.table_id} расходится с источником"))
        return ValidationReport(ok=not any(x.severity == "error" for x in issues), issues=issues, metrics={
            "fidelity_numbers": n_ok / n_total if n_total else 1.0,
            "fidelity_tables": cells_ok / cells if cells else 1.0, "stub_validator": 1.0})
