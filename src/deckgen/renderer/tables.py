"""Нативные таблицы (graphicFrame + a:tbl): заполнение ячеек строками из TableSpec без изменений,
явные кегли, ширины колонок по содержимому, числа вправо, без объединённых ячеек."""
from __future__ import annotations

from dataclasses import dataclass
from typing import Optional

from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Emu

from deckgen.contracts import BBox, TableSpec
from deckgen.fitter.measure import is_numeric, layout_table, table_grid
from deckgen.template.textstyle import CELL_MAR_LR_EMU, CELL_MAR_TB_EMU

from .ops import end_para_size, style_run


@dataclass
class TableStyle:
    style_id: Optional[str] = None      # tableStyleId из шаблона (convention: STYLE_TABLE)
    header_text: Optional[str] = None   # RRGGBB; None = из стиля таблицы
    header_fill: Optional[str] = None
    band_fill: Optional[str] = None
    border: Optional[str] = None        # горизонтальные линии между строками
    header_bold: bool = True


def _numeric_cols(grid: list[list[str]], header_n: int) -> list[bool]:
    body = grid[header_n:] or grid
    ncols = len(grid[0]) if grid else 0
    out = []
    for c in range(ncols):
        vals = [r[c] for r in body if r[c].strip()]
        out.append(bool(vals) and sum(is_numeric(v) for v in vals) >= 0.6 * len(vals))
    return out


def _tcpr(cell):
    return cell._tc.get_or_add_tcPr()


def _set_fill(cell, color: str) -> None:
    cell.fill.solid()
    from pptx.dml.color import RGBColor
    cell.fill.fore_color.rgb = RGBColor.from_string(color)


def _set_border(cell, side: str, color: Optional[str], w_emu: int = 6350) -> None:
    """side: lnL/lnR/lnT/lnB. color=None -> линия отсутствует."""
    tcPr = _tcpr(cell)
    old = tcPr.find(qn(f"a:{side}"))
    if old is not None:
        tcPr.remove(old)
    ln = etree.Element(qn(f"a:{side}"), w=str(w_emu), cap="flat", cmpd="sng", algn="ctr")
    if color:
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=color)
    else:
        etree.SubElement(ln, qn("a:noFill"))
    # порядок в tcPr: lnL, lnR, lnT, lnB, ..., fill
    order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
    idx = order.index(f"a:{side}")
    for later in order[idx + 1:] + ["a:lnTlToBr", "a:lnBlToTr", "a:cell3D", "a:noFill", "a:solidFill",
                                     "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill"]:
        nxt = tcPr.find(qn(later))
        if nxt is not None:
            nxt.addprevious(ln)
            return
    tcPr.append(ln)


def fill_table(graphic_frame, spec: TableSpec, zone: BBox, font_pt: float, style: TableStyle) -> None:
    grid, header_n = table_grid(spec)
    tl = layout_table(grid, header_n, zone.w, font_pt)
    table = graphic_frame.table
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblPr.set("firstRow", "1" if header_n else "0")
    tblPr.set("bandRow", "1")
    sid = tblPr.find(qn("a:tableStyleId"))
    if style.style_id:
        if sid is None:
            sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
        sid.text = style.style_id
    elif sid is not None:
        tblPr.remove(sid)
    for i, w in enumerate(tl.col_w):
        table.columns[i].width = Emu(w)
    for i, h in enumerate(tl.row_h):
        table.rows[i].height = Emu(h)
    graphic_frame.left, graphic_frame.top = Emu(zone.x), Emu(zone.y)
    graphic_frame.width, graphic_frame.height = Emu(sum(tl.col_w)), Emu(tl.height)

    numeric = _numeric_cols(grid, header_n)
    for r, row in enumerate(grid):
        is_hdr = r < header_n
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.margin_left = cell.margin_right = Emu(CELL_MAR_LR_EMU)
            cell.margin_top = cell.margin_bottom = Emu(CELL_MAR_TB_EMU)
            _tcpr(cell).set("anchor", "ctr")
            tf = cell.text_frame
            tf.word_wrap = True
            paras = value.split("\n") if value else [""]
            tf.text = ""  # одна пустая строка
            for pi, ptxt in enumerate(paras):
                p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
                p.text = ptxt  # \v -> <a:br/>
                p._p.get_or_add_pPr().set("algn", "r" if numeric[c] else "l")
                for run in p.runs:
                    style_run(run, font_pt, bold=True if (is_hdr and style.header_bold) else None,
                              color=style.header_text if is_hdr else None)
                end_para_size(p, font_pt)
            if is_hdr and style.header_fill:
                _set_fill(cell, style.header_fill)
            elif not is_hdr and style.band_fill and (r - header_n) % 2 == 1:
                _set_fill(cell, style.band_fill)
            if style.border:
                for side in ("lnL", "lnR"):
                    _set_border(cell, side, None)
                _set_border(cell, "lnT", style.border if r > 0 else None)
                _set_border(cell, "lnB", style.border)
