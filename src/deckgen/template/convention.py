"""ConventionAdapter: шаблон по README организаторов (макеты title*/section*/slide*/table*/last*,
фигуры-плейсхолдеры с именами *_FIELD). Главный путь — жюри подставит свой шаблон в этом формате."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.util import Emu

from deckgen.contracts import (BBox, FieldKind, FieldSpec, LayoutSpec, SlideType, StyleTokens,
                               TemplateCatalog)

from . import metrics
from .errors import TemplateError
from .ooxml import body_insets, ph_idx, resolve_font_pt, resolve_ln_spc

# Порядок важен только для читаемости; префиксы не пересекаются.
TYPE_PREFIXES: list[tuple[str, SlideType]] = [
    ("title", SlideType.title), ("section", SlideType.section), ("slide", SlideType.slide),
    ("table", SlideType.table), ("last", SlideType.last),
]
FIELD_RE = re.compile(r"(SUBTITLE|TITLE|TEXT|TABLE|IMAGE)_FIELD", re.IGNORECASE)
_KIND = {"TITLE": FieldKind.title, "SUBTITLE": FieldKind.subtitle, "TEXT": FieldKind.text,
         "TABLE": FieldKind.table, "IMAGE": FieldKind.image}
SYNTH_SUBTITLE = "SUBTITLE_FIELD"   # синтезированное поле (нет в макете) -> рендерер ставит TextBox


def slide_type_of(layout_name: str) -> Optional[SlideType]:
    n = layout_name.strip().lower()
    for prefix, t in TYPE_PREFIXES:
        if n.startswith(prefix):
            return t
    return None


def field_base(shape_name: str) -> Optional[str]:
    m = FIELD_RE.search(shape_name or "")
    return m.group(1).upper() if m else None


def iter_layouts(prs):
    """Все макеты всех мастеров с уникальными id (дубли имён получают суффикс #N)."""
    seen: dict[str, int] = {}
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            n = seen.get(layout.name, 0) + 1
            seen[layout.name] = n
            yield (layout.name if n == 1 else f"{layout.name}#{n}"), layout, master


def find_layout(prs, layout_id: str):
    for lid, layout, _ in iter_layouts(prs):
        if lid == layout_id:
            return layout
    raise TemplateError(f"Макет '{layout_id}' не найден в шаблоне")


def is_convention(prs) -> bool:
    for lid, layout, _ in iter_layouts(prs):
        if slide_type_of(layout.name) and any(field_base(s.name) == "TITLE" for s in layout.shapes):
            return True
    return False


def open_presentation(path: str):
    p = Path(path)
    if not p.is_file():
        raise TemplateError(f"Шаблон не найден: {path}")
    try:
        return Presentation(str(p))
    except Exception as e:  # битый zip, не pptx и т.п.
        raise TemplateError(f"Не удалось открыть шаблон {path} как PPTX: {e}") from e


class ConventionAdapter:
    def profile(self, template_path: str) -> TemplateCatalog:
        prs = open_presentation(template_path)
        if not is_convention(prs):
            names = [lid for lid, _, _ in iter_layouts(prs)]
            raise TemplateError(
                "Шаблон не соответствует README организаторов: нужны макеты с именами "
                "title*/section*/slide*/table*/last* и фигурой TITLE_FIELD. "
                f"Найдены макеты: {names}")
        sw, sh = int(prs.slide_width), int(prs.slide_height)
        layouts: list[LayoutSpec] = []
        for lid, layout, master in iter_layouts(prs):
            st = slide_type_of(layout.name)
            if st is None:
                continue
            fields = self._fields(layout, master._element, st, sh)
            if not any(f.kind == FieldKind.title for f in fields):
                continue
            layouts.append(LayoutSpec(layout_id=lid, slide_type=st, fields=fields, source="layout"))
        style = self._style(layouts)
        return TemplateCatalog(template_id=Path(template_path).stem, adapter="convention",
                               slide_w=sw, slide_h=sh, layouts=layouts, style=style)

    # --- поля ---
    def _fields(self, layout, master_el, st: SlideType, slide_h: int) -> list[FieldSpec]:
        fields: list[FieldSpec] = []
        counts: dict[str, int] = {}
        for shape in layout.shapes:
            base = field_base(shape.name)
            if base is None:
                continue
            kind = _KIND[base]
            counts[base] = counts.get(base, 0) + 1
            name = f"{base}_FIELD" if counts[base] == 1 else f"{base}_FIELD#{counts[base]}"
            sp = shape._element
            idx = ph_idx(sp) if shape.is_placeholder else None
            bbox = BBox(x=int(shape.left or 0), y=int(shape.top or 0), w=int(shape.width or 0), h=int(shape.height or 0))
            font_pt = resolve_font_pt(sp, master_el) if kind != FieldKind.image else 0
            max_chars = max_lines = 0
            if kind in (FieldKind.title, FieldKind.subtitle, FieldKind.text):
                max_chars, max_lines = metrics.capacity(bbox.w, bbox.h, font_pt,
                                                        resolve_ln_spc(sp, master_el), body_insets(sp))
            fields.append(FieldSpec(name=name, kind=kind, bbox=bbox, font_pt=round(font_pt, 2),
                                    max_chars=max_chars, max_lines=max_lines, placeholder_idx=idx,
                                    required=kind in (FieldKind.title, FieldKind.text, FieldKind.table)))
        if st in (SlideType.title, SlideType.section) and not any(f.kind == FieldKind.subtitle for f in fields):
            synth = self._synth_subtitle(layout, fields, slide_h)
            if synth:
                fields.append(synth)
        return fields

    @staticmethod
    def _synth_subtitle(layout, fields: list[FieldSpec], slide_h: int) -> Optional[FieldSpec]:
        """В макете нет SUBTITLE_FIELD: берём свободную полосу под заголовком до ближайшей
        декоративной фигуры (картинка/полоса внизу), рендерер поставит туда TextBox."""
        title = next(f for f in fields if f.kind == FieldKind.title)
        y0 = title.bbox.y + title.bbox.h + Emu(91440)
        y1 = int(slide_h * 0.92)
        for s in layout.shapes:
            if s.is_placeholder or s.top is None or s.width is None:
                continue
            overlaps_x = s.left < title.bbox.x + title.bbox.w and s.left + s.width > title.bbox.x
            if overlaps_x and s.top >= y0 and s.top < y1:
                y1 = int(s.top)
        h = y1 - y0 - 91440
        if h < Emu(914400 * 0.4):
            return None
        font = max(14.0, min(20.0, round(title.font_pt * 0.45)))
        bbox = BBox(x=title.bbox.x, y=int(y0), w=title.bbox.w, h=int(h))
        mc, ml = metrics.capacity(bbox.w, bbox.h, font)
        return FieldSpec(name=SYNTH_SUBTITLE, kind=FieldKind.subtitle, bbox=bbox, font_pt=font,
                         max_chars=mc, max_lines=ml, placeholder_idx=None, required=False)

    @staticmethod
    def _style(layouts: list[LayoutSpec]) -> StyleTokens:
        def first(kind: FieldKind, types: tuple[SlideType, ...]) -> Optional[float]:
            for l in layouts:
                if l.slide_type in types:
                    for f in l.fields:
                        if f.kind == kind and f.font_pt:
                            return f.font_pt
            return None
        content = (SlideType.slide, SlideType.table)
        title_pt = first(FieldKind.title, content) or 20
        body = first(FieldKind.text, content) or 14
        table = first(FieldKind.table, content) or body
        return StyleTokens(font_family="Liberation Sans", title_pt=title_pt,
                           body_pt_max=body, body_pt_min=min(10.0, body),
                           table_pt_max=min(12.0, table), table_pt_min=min(8.0, table))
