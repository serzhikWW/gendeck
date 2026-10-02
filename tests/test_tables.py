"""ML2-5: качество таблиц — нарезка 12×6 и 30×8, ячейки как в источнике, стиль шаблона в ячейках."""
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

from deckgen.contracts import DeckPlan, IngestResult, SlideSpec, SlideType, SourceTable, TableSpec
from deckgen.fitter import DeterministicFitter
from deckgen.renderer import PptxRenderer
from deckgen.renderer.tablestyle import resolve_color
from deckgen.template import profile

ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = {"example": str(ROOT / "templates" / "example_fields.pptx"),
             "interrao": str(ROOT / "templates" / "interrao_corporate.pptx")}


def _table(n_rows: int, n_cols: int) -> SourceTable:
    cols = ["Показатель"] + [f"{2018 + i}" for i in range(n_cols - 1)]
    rows = [[f"Параметр энергосистемы {r + 1}"] + [f"{(r + 1) * (c + 3)},{c}" if c % 3 else "≥ 10" if c else "—"
                                                    for c in range(n_cols - 1)] for r in range(n_rows)]
    return SourceTable(table_id="t", columns=cols, rows=rows)


def _render(tpl_key: str, st: SourceTable, tmp_path, bullets=()):
    cat = profile(TEMPLATES[tpl_key])
    src = IngestResult(raw_text="", source_hash="h", sections=[], tables=[st])
    plan = DeckPlan(deck_title="x", slides=[SlideSpec(type=SlideType.table, title="Баланс мощности",
                                                      bullets=list(bullets), table=TableSpec(source_id="t"))])
    deck = DeterministicFitter().fit(plan, src, cat)
    out = tmp_path / f"{tpl_key}.pptx"
    PptxRenderer().render(deck, TEMPLATES[tpl_key], cat, str(out))
    return deck, Presentation(str(out))


def _tables(prs):
    return [sh for s in prs.slides for sh in s.shapes if sh.has_table]


@pytest.mark.parametrize("tpl", list(TEMPLATES))
def test_12x6_one_slide(tpl, tmp_path):
    st = _table(11, 6)  # 12 строк с шапкой
    deck, prs = _render(tpl, st, tmp_path)
    assert len(deck.slides) == 1
    t = _tables(prs)[0].table
    assert [[c.text for c in r.cells] for r in t.rows] == [st.columns] + st.rows


@pytest.mark.parametrize("tpl", list(TEMPLATES))
def test_30x8_split_with_header(tpl, tmp_path):
    st = _table(29, 8)  # 30 строк с шапкой
    deck, prs = _render(tpl, st, tmp_path)
    tables = _tables(prs)
    assert len(tables) >= 2 and len(tables) == len(deck.slides)
    body = []
    for gf in tables:
        rows = [[c.text for c in r.cells] for r in gf.table.rows]
        assert rows[0] == st.columns  # шапка повторяется
        body += rows[1:]
        slide_h = prs.slide_height
        assert gf.top + gf.height <= slide_h  # таблица не вылезает за слайд
    assert body == st.rows
    titles = [s.spec.title for s in deck.slides]
    assert titles[0].endswith(f"(1/{len(titles)})")


def test_text_plus_big_table_goes_to_separate_slides(tmp_path):
    deck, _ = _render("example", _table(29, 8), tmp_path, bullets=["Пиковая нагрузка выросла на 3,1%"])
    assert deck.slides[0].spec.table is None and deck.slides[0].spec.bullets
    assert all(s.spec.table is not None for s in deck.slides[1:])


def test_alignment_numbers_right_text_left(tmp_path):
    _, prs = _render("example", _table(5, 4), tmp_path)
    t = _tables(prs)[0].table
    assert t.cell(1, 0).text_frame.paragraphs[0].alignment == PP_ALIGN.LEFT
    for c in (1, 2, 3):
        assert t.cell(1, c).text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT


def test_no_merged_cells_and_explicit_font(tmp_path):
    for tpl in TEMPLATES:
        _, prs = _render(tpl, _table(5, 4), tmp_path)
        tbl = _tables(prs)[0]._element
        for tc in tbl.iter(qn("a:tc")):
            assert not (set(tc.attrib) & {"gridSpan", "rowSpan", "hMerge", "vMerge"})
        for r in tbl.iter(qn("a:r")):
            rpr = r.find(qn("a:rPr"))
            assert rpr is not None and rpr.get("sz")
            assert rpr.find(qn("a:latin")).get("typeface") == "Liberation Sans"


def test_template_style_baked_into_cells(tmp_path):
    """example_fields: STYLE_TABLE = стиль с рамками accent2 (ED7D31) и зеброй accent2 20% — в ячейках явно."""
    _, prs = _render("example", _table(5, 4), tmp_path)
    tbl = _tables(prs)[0]._element
    assert tbl.find(f".//{qn('a:tableStyleId')}").text == "{5DA37D80-6434-44D0-A028-1B22A696006F}"
    xml = tbl.xml if hasattr(tbl, "xml") else ""
    from lxml import etree
    xml = etree.tostring(tbl, encoding="unicode")
    assert 'val="ED7D31"' in xml and 'val="FBE5D6"' in xml
    # первая строка шапки — жирная
    hdr = _tables(prs)[0].table.cell(0, 0).text_frame.paragraphs[0].runs[0]
    assert hdr.font.bold


def test_interrao_corporate_table_style(tmp_path):
    _, prs = _render("interrao", _table(5, 4), tmp_path)
    t = _tables(prs)[0].table
    hdr = t.cell(0, 1).text_frame.paragraphs[0].runs[0]
    assert str(hdr.font.color.rgb) == "F18400" and hdr.font.bold


def test_resolve_color_transforms():
    from lxml import etree
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    theme = {"accent2": "ED7D31", "lt1": "FFFFFF"}
    el = etree.fromstring(f'<a:schemeClr xmlns:a="{A}" val="accent2"><a:alpha val="20000"/></a:schemeClr>')
    assert resolve_color(el, theme) == "FBE5D6"
    el = etree.fromstring(f'<a:schemeClr xmlns:a="{A}" val="bg1"><a:lumMod val="85000"/></a:schemeClr>')
    assert resolve_color(el, theme) == "D9D9D9"
