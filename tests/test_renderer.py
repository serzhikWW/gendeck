import re
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from deckgen.fitter import DeterministicFitter
from deckgen.renderer import PptxRenderer
from deckgen.template import profile

import ml2_fixtures as fx

ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = {"example": ROOT / "templates" / "example_fields.pptx",
             "interrao": ROOT / "templates" / "interrao_corporate.pptx"}


@pytest.fixture(scope="module", params=list(TEMPLATES))
def rendered(request, tmp_path_factory):
    tpl = str(TEMPLATES[request.param])
    cat = profile(tpl)
    deck = DeterministicFitter().fit(fx.plan(), fx.ingest(), cat)
    out = tmp_path_factory.mktemp(request.param) / "out.pptx"
    PptxRenderer().render(deck, tpl, cat, str(out))
    return deck, out


def test_slide_count_and_titles(rendered):
    deck, out = rendered
    prs = Presentation(str(out))
    assert len(prs.slides) == len(deck.slides)  # слайды-образцы шаблона удалены
    texts = ["\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
    for fs, t in zip(deck.slides, texts):
        if fs.spec.type.value != "last" or "Спасибо" not in t.upper().title():
            assert fs.spec.title in t or fs.spec.type.value == "last"


def test_tables_native_and_verbatim(rendered):
    deck, out = rendered
    prs = Presentation(str(out))
    for fs, s in zip(deck.slides, prs.slides):
        tables = [sh for sh in s.shapes if sh.has_table]
        if fs.spec.table is None:
            assert not tables
            continue
        assert len(tables) == 1
        tbl = tables[0].table
        got = [[c.text for c in r.cells] for r in tbl.rows]
        exp = ([fs.spec.table.columns] if fs.spec.table.columns else []) + fs.spec.table.rows
        assert got == exp
        assert tables[0]._element.find(f".//{qn('a:tbl')}") is not None


def test_no_autofit_explicit_sizes_liberation(rendered):
    _, out = rendered
    with zipfile.ZipFile(out) as z:
        for name in z.namelist():
            if re.match(r"ppt/(slides|slideLayouts|slideMasters)/[^/]+\.xml$", name):
                xml = z.read(name).decode("utf8")
                assert "normAutofit" not in xml and "spAutoFit" not in xml, name
    prs = Presentation(str(out))
    for s in prs.slides:
        for sh in s.shapes:
            frames = [sh.text_frame] if sh.has_text_frame else []
            if sh.has_table:
                frames += [c.text_frame for r in sh.table.rows for c in r.cells]
            for tf in frames:
                for p in tf.paragraphs:
                    for r in p.runs:
                        if r.text.strip():
                            assert r.font.size is not None, r.text
                            # фигуры, которые мы заполняли, — явно Liberation Sans
                            assert r.font.name in (None, "Liberation Sans"), (r.font.name, r.text)


def test_no_empty_placeholders(rendered):
    _, out = rendered
    for s in Presentation(str(out)).slides:
        for sh in s.placeholders:
            el = sh._element
            if el.tag == qn("p:sp"):
                assert sh.text_frame.text.strip() or el.find(f".//{qn('a:fld')}") is not None


def test_package_has_no_dangling_parts(rendered):
    _, out = rendered
    with zipfile.ZipFile(out) as z:
        names = set(z.namelist())
        assert z.testzip() is None
        slides = sorted((n for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                        key=lambda n: int(re.findall(r"\d+", n)[-1]))
        assert slides == [f"ppt/slides/slide{i}.xml" for i in range(1, len(slides) + 1)]
        referenced = set()
        for n in names:
            if n.endswith(".rels"):
                base = n.replace("_rels/", "").replace(".rels", "")
                base_dir = base.rsplit("/", 1)[0] if "/" in base else ""
                for tgt, mode in re.findall(r'Target="([^"]+)"(?:\s+TargetMode="(\w+)")?', z.read(n).decode("utf8")):
                    if mode == "External":
                        continue
                    parts = (base_dir + "/" + tgt).split("/") if not tgt.startswith("/") else tgt[1:].split("/")
                    stack = []
                    for p in parts:
                        if p == "..":
                            stack.pop()
                        elif p and p != ".":
                            stack.append(p)
                    path = "/".join(stack)
                    assert path in names, f"битая ссылка {n} -> {tgt}"
                    referenced.add(path)
        media = {n for n in names if n.startswith("ppt/media/")}
        assert media <= referenced, f"неиспользуемые media: {media - referenced}"


def test_notes_written(rendered):
    deck, out = rendered
    prs = Presentation(str(out))
    for fs, s in zip(deck.slides, prs.slides):
        if fs.spec.notes:
            assert s.notes_slide.notes_text_frame.text == fs.spec.notes
