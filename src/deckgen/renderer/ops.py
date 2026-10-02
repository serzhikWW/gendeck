"""Операции над пакетом PPTX, которых нет в python-pptx: клон/удаление слайдов, перенумерация,
чистка autofit и пустых плейсхолдеров, стиль шрифта ранов."""
from __future__ import annotations

import copy

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Pt

FONT = "Liberation Sans"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_GRAPHIC_TABLE = "http://schemas.openxmlformats.org/drawingml/2006/table"


# ---------- слайды ----------
def clone_slide(prs, src):
    """Новый слайд в конце: тот же макет, копия фигур и фона образца. Диаграммы/OLE не копируются
    (их части нельзя делить между слайдами, а для генерации они не нужны)."""
    new = prs.slides.add_slide(src.slide_layout)
    tree = new.shapes._spTree
    for el in list(tree):
        if el.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr"), qn("p:extLst")):
            tree.remove(el)
    anchor = tree.find(qn("p:extLst"))
    for el in src.shapes._spTree:
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr"), qn("p:extLst")):
            continue
        if _is_chart_or_ole(el):
            continue
        dup = copy.deepcopy(el)
        _remap_rels(dup, src.part, new.part)
        if anchor is not None:
            anchor.addprevious(dup)
        else:
            tree.append(dup)
    src_bg = src._element.cSld.find(qn("p:bg"))
    if src_bg is not None:
        bg = copy.deepcopy(src_bg)
        _remap_rels(bg, src.part, new.part)
        new._element.cSld.insert(0, bg)
    return new


def _is_chart_or_ole(el) -> bool:
    if el.tag == qn("p:graphicFrame"):
        gd = el.find(f".//{qn('a:graphicData')}")
        return gd is None or gd.get("uri") != _GRAPHIC_TABLE
    if el.tag == qn("p:grpSp"):
        return any(_is_chart_or_ole(c) for c in el.iter(qn("p:graphicFrame")))
    return False


def _remap_rels(el, src_part, dst_part) -> None:
    for node in el.iter():
        for attr, val in list(node.attrib.items()):
            if not attr.startswith(f"{{{_R_NS}}}"):
                continue
            rel = src_part.rels.get(val)
            if rel is None:
                continue
            if rel.is_external:
                new_id = dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                new_id = dst_part.relate_to(rel.target_part, rel.reltype)
            node.set(attr, new_id)


def delete_slides(prs, slides) -> None:
    doomed = {id(s.part) for s in slides}
    lst = prs.slides._sldIdLst
    for sld_id in list(lst):
        if id(prs.part.related_part(sld_id.rId)) in doomed:
            rid = sld_id.rId
            lst.remove(sld_id)
            prs.part.drop_rel(rid)


def renumber(prs) -> None:
    """slideN.xml / notesSlideN.xml по порядку (недостижимые старые части в файл не попадут)."""
    for i, s in enumerate(prs.slides, 1):
        s.part.partname = PackURI(f"/ppt/slides/slide{i}.xml")
        if s.has_notes_slide:
            s.notes_slide.part.partname = PackURI(f"/ppt/notesSlides/notesSlide{i}.xml")


# ---------- текст ----------
def no_autofit(txBody) -> None:
    """Явно запрещаем автоподбор (normAutofit/spAutoFit в Р7 ведут себя непредсказуемо)."""
    bp = txBody.find(qn("a:bodyPr"))
    if bp is None:
        bp = etree.SubElement(txBody, qn("a:bodyPr"))
        txBody.insert(0, bp)
    for tag in ("a:normAutofit", "a:spAutoFit", "a:noAutofit"):
        for el in bp.findall(qn(tag)):
            bp.remove(el)
    na = etree.SubElement(bp, qn("a:noAutofit"))
    # noAutofit должен идти до scene3d/sp3d/flatTx/extLst
    for tag in ("a:scene3d", "a:sp3d", "a:flatTx", "a:extLst"):
        nxt = bp.find(qn(tag))
        if nxt is not None:
            nxt.addprevious(na)
            break


def style_run(run, size_pt: float, bold=None, color: str | None = None) -> None:
    f = run.font
    f.size = Pt(size_pt)
    f.name = FONT
    if bold is not None:
        f.bold = bold
    if color:
        f.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    latin = rpr.find(qn("a:latin"))
    after = latin
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
            after.addnext(el)
        el.set("typeface", FONT)
        after = el


def end_para_size(p, size_pt: float) -> None:
    epr = p._p.find(qn("a:endParaRPr"))
    if epr is None:
        epr = etree.SubElement(p._p, qn("a:endParaRPr"))
    epr.set("sz", str(int(round(size_pt * 100))))


# ---------- финальная чистка ----------
def is_empty_placeholder(shape) -> bool:
    if not shape.is_placeholder:
        return False
    el = shape._element
    if el.tag == qn("p:sp"):
        if el.find(f".//{qn('a:fld')}") is not None:
            return False
        txb = el.find(qn("p:txBody"))
        return txb is None or not "".join(t.text or "" for t in txb.iter(qn("a:t"))).strip()
    if el.tag == qn("p:pic"):
        blip = el.find(f".//{qn('a:blip')}")
        return blip is None or not blip.get(qn("r:embed"))
    return False


def remove_empty_placeholders(slide) -> int:
    n = 0
    for shape in list(slide.shapes):
        if is_empty_placeholder(shape):
            shape._element.getparent().remove(shape._element)
            n += 1
    return n




def strip_autofit_everywhere(prs) -> None:
    """Заменить normAutofit/spAutoFit на noAutofit в слайдах, макетах и мастерах."""
    parts = [s.part for s in prs.slides]
    for m in prs.slide_masters:
        parts.append(m.part)
        parts.extend(l.part for l in m.slide_layouts)
    for part in parts:
        root = part._element
        for tag in ("a:normAutofit", "a:spAutoFit"):
            for el in list(root.iter(qn(tag))):
                el.addprevious(etree.Element(qn("a:noAutofit")))
                el.getparent().remove(el)
